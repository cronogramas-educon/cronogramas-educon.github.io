#!/usr/bin/env python3
"""Genera el Excel de asistencia de un curso.

  python scripts/generar_asistencia_xlsx.py --plantilla SALIDA.xlsx      plantilla en blanco para copiar a cada curso nuevo
  python scripts/generar_asistencia_xlsx.py DATOS.json SALIDA.xlsx       un curso con sus datos
DATOS.json: {"curso": "...", "minimo": 0.8, "clases": [{"fecha": "2026-10-01", "situacion": "", "nueva": ""}, ...],
             "estudiantes": [{"documento": "", "nombre": "...", "marcas": [1, 0, null, ...]}, ...]}
Los datos de estudiantes no viven en el repositorio (es público): el JSON se escribe fuera de él.

El archivo cabe hasta 40 clases. Cuántas tiene el curso lo dice la tabla Parametros de la hoja Resumen (fila "Clases del curso"):
se escribe a mano (el sitio ya la toma del Cuadro de horas; aquí solo afecta los colores del archivo). Las columnas de clases que sobran quedan "Sin programar".
Hojas: Instrucciones, Asistencia (tabla Table1, una columna por clase), Resumen (tabla Parametros y cuentas) y Cronograma (fechas y situación de cada clase).
Las cuentas son fórmulas: el sitio lee las columnas calculadas (scripts/asistencia.py)."""
import json, sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as colletra
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableFormula, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parent))
from generar_registro_xlsx import AZUL, AZUL_CLARO, FUENTE, GRIS, ROJO, ROSA, _paso, _seccion, _texto

MAXC = 40            # clases que caben en el archivo
PRIMERA = 7          # primera fila de estudiantes (el encabezado de la tabla es la 6)
HOLGURA = 200        # filas que las cuentas y las reglas cubren, para poder agregar estudiantes
COL0 = 4             # la clase 1 está en la columna D
CALCULADAS = ["Asistió", "Faltas", "Sin marcar", "Faltas que le quedan", "% hasta hoy", "Puede llegar hasta", "Estado", "Clases programadas", "Clases del curso"]
SEGUIMIENTO = ["Contactado", "En conversación", "Resuelto"]
ESTADOS = {"Al día": ("E2F0D9", "375623"), "Atención": ("FFF2CC", "7F6000"), "En riesgo": ("FBE0D6", "C00000"),
           "Perdió el mínimo": ("C00000", "FFFFFF"), "Aprobó": ("DDEBF7", "1F3864"), "Falta configurar": ("F2F2F2", "7F7F7F")}
# filas de la hoja Resumen: la tabla Parametros va de la 4 (encabezado) a la 8; las cuentas, debajo
R = {"clases": 5, "minimo": 6, "atencion": 7, "riesgo": 8,
     "programadas": 11, "dictadas": 12, "pendientes": 13, "mimas": 14, "permitidas": 15, "estudiantes": 16, "promedio": 17, "sinmarcar": 26}
RES = lambda k: f"Resumen!$C${R[k]}"
ULT_CL = colletra(COL0 + MAXC - 1)     # última columna de clases (AQ)
COL = {nombre: colletra(COL0 + MAXC + k) for k, nombre in enumerate(CALCULADAS + ["Seguimiento"])}
TOTAL_COLS = COL0 + MAXC + len(CALCULADAS)


def _celda(ws, ref, valor, **kw):
    c = ws[ref]
    c.value = valor
    c.font = Font(name=FUENTE, size=kw.pop("size", 11), bold=kw.pop("bold", False), color=kw.pop("color", "000000"))
    if kw.get("fill"):
        c.fill = PatternFill("solid", fgColor=kw["fill"])
    kw.pop("fill", None)
    c.alignment = Alignment(horizontal=kw.pop("h", None), vertical="center", wrap_text=kw.pop("wrap", False), indent=kw.pop("indent", 0))
    if kw.get("fmt"):
        c.number_format = kw["fmt"]
    return c


def _banner(ws, titulo, hasta):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=hasta)
    _celda(ws, "A1", titulo, size=20, bold=True, color="FFFFFF", indent=1)
    for col in range(1, hasta + 1):
        ws.cell(1, col).fill = PatternFill("solid", fgColor=AZUL)
        ws.cell(2, col).fill = PatternFill("solid", fgColor=ROJO)
    ws.row_dimensions[1].height = 40
    ws.row_dimensions[2].height = 5


def _hoja_cronograma(wb, clases):
    ws = wb.create_sheet("Cronograma")
    ws.sheet_view.showGridLines = False
    _banner(ws, "Cronograma de clases", 6)
    for col, w in zip("ABCDEF", (12, 14, 14, 18, 16, 54)):
        ws.column_dimensions[col].width = w
    _celda(ws, "A3", "Opcional. Las fechas son solo para verlas en la hoja Asistencia: el sitio toma las fechas del Cuadro de horas. Cambia la Situación solo si una clase se reprograma o se cancela.", color="595959")
    for j, h in enumerate(("Clase", "Fecha", "Día", "Situación", "Nueva fecha", "Nota"), 1):
        _celda(ws, f"{colletra(j)}5", h, bold=True, color="FFFFFF", fill=AZUL, h="center")
    ws.row_dimensions[5].height = 24
    for i in range(1, MAXC + 1):
        r = 5 + i
        c = clases[i - 1] if i <= len(clases) else {}
        _celda(ws, f"A{r}", f"Clase {i}", bold=True)
        _celda(ws, f"B{r}", date.fromisoformat(c["fecha"]) if c.get("fecha") else None, fmt="dd/mm/yyyy", h="center")
        _celda(ws, f"C{r}", f'=IF(B{r}="","",CHOOSE(WEEKDAY(B{r},2),"lunes","martes","miércoles","jueves","viernes","sábado","domingo"))', h="center")
        _celda(ws, f"D{r}", c.get("situacion") or None, h="center")
        _celda(ws, f"E{r}", date.fromisoformat(c["nueva"]) if c.get("nueva") else None, fmt="dd/mm/yyyy", h="center")
        _celda(ws, f"F{r}", c.get("nota") or None)
        for col in "ABCDEF":
            ws[f"{col}{r}"].border = Border(bottom=Side(style="thin", color="BFBFBF"))
    ult = 5 + MAXC
    dv = DataValidation(type="list", formula1='"Reprogramada,Cancelada"', allow_blank=True, showErrorMessage=True, showInputMessage=True)
    dv.promptTitle, dv.prompt = "Situación", "Déjala vacía si la clase sigue igual. Elige Reprogramada o Cancelada."
    dv.error, dv.errorTitle = "Elige Reprogramada o Cancelada, o deja la celda vacía.", "Situación"
    ws.add_data_validation(dv)
    dv.add(f"D6:D{ult}")
    ws.freeze_panes = "A6"


def _hoja_resumen(wb, n_clases, minimo):
    ws = wb.create_sheet("Resumen")
    ws.sheet_view.showGridLines = False
    _banner(ws, "Resumen de asistencia", 4)
    for col, w in zip("ABCD", (2, 52, 22, 70)):
        ws.column_dimensions[col].width = w
    est = f"Asistencia!$D$4:${ULT_CL}$4"
    rango = lambda nombre: f"Asistencia!${COL[nombre]}${PRIMERA}:${COL[nombre]}${PRIMERA + HOLGURA}"

    def seccion(r, texto):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        _celda(ws, f"B{r}", texto, bold=True, size=13, color=AZUL)
        for cc in range(2, 5):
            ws.cell(r, cc).border = Border(bottom=Side(style="medium", color=ROJO))
        ws.row_dimensions[r].height = 26

    def cuenta(k, etiqueta, formula, nota="", fmt=None):
        r = R[k]
        _celda(ws, f"B{r}", etiqueta, indent=1)
        _celda(ws, f"C{r}", formula, bold=True, h="center", fill=GRIS, color=AZUL, fmt=fmt)
        _celda(ws, f"D{r}", nota, color="595959", wrap=True)
        ws.row_dimensions[r].height = 30 if len(nota) > 80 else 22

    seccion(3, "Datos de este archivo (las celdas amarillas se pueden cambiar)")
    # tabla Parametros: un flujo de Power Automate puede actualizar sus filas por el nombre del dato
    datos = [("clases", "Clases del curso", n_clases, "Cuántas clases tiene el curso en el Cuadro de horas (el número de la última clase). Máximo 40. El sitio ya lo toma del Cuadro: aquí solo sirve para los colores y el Estado de este archivo.", "0"),
             ("minimo", "Asistencia mínima", minimo, "Porcentaje mínimo de clases para certificar.", "0%"),
             ("atencion", "Aviso de atención", 0.5, "Marca Atención cuando use este % de las faltas permitidas. Con 7 permitidas, empieza en la falta 4.", "0%"),
             ("riesgo", "Aviso de riesgo", 0.8, "Marca En riesgo cuando use este % de las faltas permitidas. Con 7 permitidas, empieza en la falta 6 (queda una).", "0%")]
    for j, h in enumerate(("Dato", "Valor", "Nota"), 2):
        _celda(ws, f"{colletra(j)}4", h, bold=True, color="FFFFFF", fill=AZUL, h="center" if j == 3 else None, indent=1 if j != 3 else 0)
    for k, etiqueta, valor, nota, fmt in datos:
        r = R[k]
        _celda(ws, f"B{r}", etiqueta, indent=1)
        _celda(ws, f"C{r}", valor, bold=True, h="center", fill="FFF2CC", color=AZUL, fmt=fmt)
        _celda(ws, f"D{r}", nota, color="595959", wrap=True)
        ws.row_dimensions[r].height = 44 if len(nota) > 150 else 30 if len(nota) > 80 else 22
    t = Table(displayName="Parametros", ref=f"B4:D{R['riesgo']}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(MAXC), allow_blank=True, showErrorMessage=True, showInputMessage=True)
    dv.error, dv.errorTitle = f"Escribe un número de 1 a {MAXC}.", "Clases del curso"
    dv.promptTitle, dv.prompt = "Clases del curso", f"Número de clases, de 1 a {MAXC}."
    ws.add_data_validation(dv)
    dv.add(f"C{R['clases']}")
    ws.conditional_formatting.add(f"C{R['clases']}", FormulaRule(formula=[f'C{R["clases"]}=""'], fill=PatternFill("solid", bgColor=ROSA, fgColor=ROSA)))

    seccion(10, "Cuentas automáticas (no se editan)")
    cuenta("programadas", "Sesiones programadas", f'=IF(N(C{R["clases"]})=0,0,C{R["clases"]}-COUNTIF({est},"Cancelada"))', "Las clases del curso menos las canceladas.")
    cuenta("dictadas", "Sesiones dictadas", f'=COUNTIF({est},"Dictada")', "Una clase cuenta como dictada cuando tiene alguna marca en la hoja Asistencia.")
    cuenta("pendientes", "Sesiones por dictar", f"=C{R['programadas']}-C{R['dictadas']}")
    cuenta("mimas", "Asistencias mínimas para certificar", f"=ROUNDUP(ROUND(C{R['programadas']}*C{R['minimo']},6),0)", "Se redondea hacia arriba.")
    cuenta("permitidas", "Faltas permitidas", f"=C{R['programadas']}-C{R['mimas']}", "Con más faltas que estas ya no se llega al mínimo.")
    cuenta("estudiantes", "Estudiantes", f'=COUNTA(Asistencia!$C${PRIMERA}:$C${PRIMERA + HOLGURA})')
    cuenta("promedio", "Asistencia promedio hasta hoy", f'=IFERROR(SUM({rango("Asistió")})/(SUM({rango("Asistió")})+SUM({rango("Faltas")})),"")', "", "0.0%")
    seccion(19, "Estudiantes por estado")
    for k, nombre in enumerate(("Al día", "Atención", "En riesgo", "Perdió el mínimo", "Aprobó")):
        fondo, letra = ESTADOS[nombre]
        r = 20 + k
        _celda(ws, f"B{r}", nombre, bold=True, color=letra, fill=fondo, indent=1)
        _celda(ws, f"C{r}", f'=COUNTIF({rango("Estado")},"{nombre}")', bold=True, h="center", color=AZUL, fill=GRIS)
    _celda(ws, f"B{R['sinmarcar']}", "Casillas sin marcar en clases ya dictadas", indent=1)
    _celda(ws, f"C{R['sinmarcar']}", f'=SUM({rango("Sin marcar")})', bold=True, h="center", color=AZUL, fill=GRIS)
    _celda(ws, f"D{R['sinmarcar']}", "Debería ser 0: si no, falta marcar 1 o 0 a alguien.", color="595959")


def _hoja_asistencia(wb, curso, estudiantes, hay_clases):
    ws = wb.active
    ws.title = "Asistencia"
    hay = max(len(estudiantes) + 6, 12)           # filas con fórmulas: los estudiantes y algunas libres
    ult = PRIMERA + hay - 1
    _banner(ws, "Asistencia", 14)
    _celda(ws, "A3", "Cada clase es una columna.  1 = asistió (verde)  ·  0 = faltó (rojo)  ·  vacío = aún no se marca. Una clase sin ninguna marca todavía no se dictó. Las columnas grises se calculan solas: no las escribas.", color="595959")
    _celda(ws, "C4", "Estado de la sesión", bold=True, h="right", color=AZUL)
    _celda(ws, "C5", "Fecha", bold=True, h="right", color=AZUL)
    cab = ["N°", "Documento", "Nombre del estudiante"] + [f"Clase {i}" for i in range(1, MAXC + 1)] + CALCULADAS + ["Seguimiento"]
    for j, h in enumerate(cab, 1):
        marca = j <= COL0 + MAXC - 1
        _celda(ws, f"{colletra(j)}6", h, bold=True, h="center", wrap=True, color="FFFFFF" if marca else ("1F3864" if h == "Seguimiento" else "595959"),
               fill=AZUL if marca else (AZUL_CLARO if h == "Seguimiento" else "D9D9D9"))
        ws.column_dimensions[colletra(j)].width = 6 if j == 1 else (16 if j == 2 else (42 if j == 3 else (9 if marca else 14)))
    ws.row_dimensions[6].height = 34
    ws.column_dimensions[COL["Estado"]].width = 18
    ws.column_dimensions[COL["Seguimiento"]].width = 20
    for i in range(1, MAXC + 1):
        L, rc = colletra(COL0 + i - 1), 5 + i
        _celda(ws, f"{L}4", f'=IF({i}>{RES("clases")},"Sin programar",IF(Cronograma!$D${rc}="Cancelada","Cancelada",IF(COUNT({L}{PRIMERA}:{L}{PRIMERA + HOLGURA})>0,"Dictada",IF(Cronograma!$D${rc}="Reprogramada","Reprogramada","Pendiente"))))',
               size=9, bold=True, h="center", wrap=True)
        _celda(ws, f"{L}5", f'=IF({i}>{RES("clases")},"",IF(Cronograma!$E${rc}<>"",Cronograma!$E${rc},IF(Cronograma!$B${rc}<>"",Cronograma!$B${rc},"")))', size=10, h="center", fmt="dd/mm/yy")
    ws.row_dimensions[4].height = 28
    estr = f"$D$4:${ULT_CL}$4"
    prog, perm, pend, dic = RES("programadas"), RES("permitidas"), RES("pendientes"), RES("dictadas")
    for k in range(hay):
        r = PRIMERA + k
        e = estudiantes[k] if k < len(estudiantes) else None
        _celda(ws, f"A{r}", (k + 1) if e else None, h="center", color="595959")
        _celda(ws, f"B{r}", (e or {}).get("documento") or None, h="center", fmt="@")
        _celda(ws, f"C{r}", (e or {}).get("nombre"), bold=bool(e))
        marcas = (e or {}).get("marcas", [])
        for i in range(MAXC):
            _celda(ws, f"{colletra(COL0 + i)}{r}", marcas[i] if i < len(marcas) else None, h="center")
        fila = f"{colletra(COL0)}{r}:{ULT_CL}{r}"
        a, f_, sm, q = (f"{COL[x]}{r}" for x in ("Asistió", "Faltas", "Sin marcar", "Faltas que le quedan"))
        nom = f"$C{r}"
        uso = f'IF({perm}=0,0,ROUND({f_}/{perm},6))'
        F = {
            "Asistió": f'=IF({nom}="","",SUMPRODUCT(({estr}="Dictada")*({fila}=1)))',
            "Faltas": f'=IF({nom}="","",SUMPRODUCT(({estr}="Dictada")*({fila}=0)*({fila}<>"")))',
            "Sin marcar": f'=IF({nom}="","",{dic}-{a}-{f_})',
            "Faltas que le quedan": f'=IF({nom}="","",{perm}-{f_})',
            "% hasta hoy": f'=IF({nom}="","",IF({a}+{f_}=0,"",{a}/({a}+{f_})))',
            "Puede llegar hasta": f'=IF({nom}="","",IF({prog}=0,"",({prog}-{f_})/{prog}))',
            "Estado": (f'=IF({nom}="","",IF({prog}=0,"Falta configurar",IF({f_}>{perm},"Perdió el mínimo",IF({pend}=0,"Aprobó",'
                       f'IF({uso}>={RES("riesgo")},"En riesgo",IF({uso}>={RES("atencion")},"Atención","Al día"))))))'),
            "Clases programadas": f'=IF({nom}="","",{prog})',
            "Clases del curso": f'=IF({nom}="","",N({RES("clases")}))',
        }
        for nombre in CALCULADAS:
            _celda(ws, f"{COL[nombre]}{r}", F[nombre], h="center", fmt="0%" if nombre in ("% hasta hoy", "Puede llegar hasta") else None,
                   fill=GRIS if nombre != "Estado" else None, bold=nombre == "Estado")
        _celda(ws, f"{COL['Seguimiento']}{r}", None, h="center")
        for j in range(1, TOTAL_COLS + 1):
            ws.cell(r, j).border = Border(bottom=Side(style="thin", color="D9D9D9"))
    # tabla: las cuentas se llenan solas en las filas que se agreguen
    t = Table(displayName="Table1", ref=f"A6:{colletra(TOTAL_COLS)}{ult}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    t._initialise_columns()
    fila_t = f"Table1[[#This Row],[Clase 1]:[Clase {MAXC}]]"
    est_t = f"Asistencia!${colletra(COL0)}$4:${ULT_CL}$4"
    nom_t = "Table1[[#This Row],[Nombre del estudiante]]"
    a_t, f_t = "Table1[[#This Row],[Asistió]]", "Table1[[#This Row],[Faltas]]"
    uso_t = f'IF({perm}=0,0,ROUND({f_t}/{perm},6))'
    FT = {
        "Asistió": f'IF({nom_t}="","",SUMPRODUCT(({est_t}="Dictada")*({fila_t}=1)))',
        "Faltas": f'IF({nom_t}="","",SUMPRODUCT(({est_t}="Dictada")*({fila_t}=0)*({fila_t}<>"")))',
        "Sin marcar": f'IF({nom_t}="","",{dic}-{a_t}-{f_t})',
        "Faltas que le quedan": f'IF({nom_t}="","",{perm}-{f_t})',
        "% hasta hoy": f'IF({nom_t}="","",IF({a_t}+{f_t}=0,"",{a_t}/({a_t}+{f_t})))',
        "Puede llegar hasta": f'IF({nom_t}="","",IF({prog}=0,"",({prog}-{f_t})/{prog}))',
        "Estado": (f'IF({nom_t}="","",IF({prog}=0,"Falta configurar",IF({f_t}>{perm},"Perdió el mínimo",IF({pend}=0,"Aprobó",'
                   f'IF({uso_t}>={RES("riesgo")},"En riesgo",IF({uso_t}>={RES("atencion")},"Atención","Al día"))))))'),
        "Clases programadas": f'IF({nom_t}="","",{prog})',
        "Clases del curso": f'IF({nom_t}="","",N({RES("clases")}))',
    }
    for c, nombre in zip(t.tableColumns, cab):
        c.name = nombre
        if nombre in FT:
            c.calculatedColumnFormula = TableFormula(attr_text=FT[nombre])
    ws.add_table(t)
    # solo 1 o 0 en las marcas; el seguimiento, de una lista
    dv = DataValidation(type="list", formula1='"1,0"', allow_blank=True, showErrorMessage=True, showInputMessage=True)
    dv.error, dv.errorTitle = "Escribe 1 si asistió, 0 si faltó, o deja la casilla vacía.", "Asistencia"
    dv.promptTitle, dv.prompt = "Asistencia", "1 = asistió, 0 = faltó. Vacía = aún no se marca."
    ws.add_data_validation(dv)
    dv.add(f"{colletra(COL0)}{PRIMERA}:{ULT_CL}{PRIMERA + HOLGURA}")
    dv2 = DataValidation(type="list", formula1='"' + ",".join(SEGUIMIENTO) + '"', allow_blank=True, showErrorMessage=False, showInputMessage=True)
    dv2.promptTitle, dv2.prompt = "Seguimiento", "Anota cómo va el contacto con este estudiante. El sitio de administración lo muestra."
    ws.add_data_validation(dv2)
    dv2.add(f"{COL['Seguimiento']}{PRIMERA}:{COL['Seguimiento']}{PRIMERA + HOLGURA}")
    rango_marcas = f"{colletra(COL0)}{PRIMERA}:{ULT_CL}{PRIMERA + HOLGURA}"
    ws.conditional_formatting.add(rango_marcas, FormulaRule(formula=[f'AND({colletra(COL0)}{PRIMERA}<>"",{colletra(COL0)}{PRIMERA}=0)'], fill=PatternFill("solid", bgColor="F8CBAD", fgColor="F8CBAD"), font=Font(color="9C0006", bold=True)))
    ws.conditional_formatting.add(rango_marcas, FormulaRule(formula=[f'{colletra(COL0)}{PRIMERA}=1'], fill=PatternFill("solid", bgColor="E2F0D9", fgColor="E2F0D9")))
    for nombre, (fondo, letra) in {"Dictada": ("E2F0D9", "375623"), "Pendiente": ("F2F2F2", "7F7F7F"), "Reprogramada": ("FFF2CC", "7F6000"), "Cancelada": ("F8CBAD", "9C0006"), "Sin programar": ("FFFFFF", "BFBFBF")}.items():
        ws.conditional_formatting.add(f"{colletra(COL0)}4:{ULT_CL}4", FormulaRule(formula=[f'{colletra(COL0)}4="{nombre}"'], fill=PatternFill("solid", bgColor=fondo, fgColor=fondo), font=Font(color=letra, bold=True, size=9)))
    ws.conditional_formatting.add(f"{colletra(COL0)}6:{ULT_CL}{PRIMERA + HOLGURA}", FormulaRule(formula=[f'{colletra(COL0)}$4="Sin programar"'], fill=PatternFill("solid", bgColor="F2F2F2", fgColor="F2F2F2"), font=Font(color="BFBFBF")))
    cE = COL["Estado"]
    for nombre, (fondo, letra) in ESTADOS.items():
        ws.conditional_formatting.add(f"{cE}{PRIMERA}:{cE}{PRIMERA + HOLGURA}", FormulaRule(formula=[f'{cE}{PRIMERA}="{nombre}"'], fill=PatternFill("solid", bgColor=fondo, fgColor=fondo), font=Font(color=letra, bold=True)))
    ws.freeze_panes = f"{colletra(COL0)}{PRIMERA}"
    ws["C6"].comment = Comment("Escribe aquí el nombre completo. Puedes pegar toda la lista desde el Excel de matriculados. El resto de la fila se calcula solo.", "Asistencia", width=260, height=80)


def _hoja_instrucciones(wb, curso):
    ay = wb.create_sheet("Instrucciones")
    ay.sheet_view.showGridLines = False
    for c, w in zip("ABCDE", (2, 24, 16, 64, 44)):
        ay.column_dimensions[c].width = w
    ay.merge_cells("B1:E1")
    _celda(ay, "B1", "Control de asistencia", size=22, bold=True, color="FFFFFF", indent=1)
    for c in range(2, 6):
        ay.cell(1, c).fill = PatternFill("solid", fgColor=AZUL)
        ay.cell(2, c).fill = PatternFill("solid", fgColor=ROJO)
    ay.row_dimensions[1].height = 48
    ay.row_dimensions[2].height = 5
    _texto(ay, 3, 2, 5, "Este archivo guarda una sola cosa: quién asistió a cada clase. Todo lo demás (fechas, módulos, profesores, número de clases) sale del Cuadro de horas del curso, y el sitio de administración calcula los avisos solo.", 44, negrita=True, color=AZUL)
    f = _seccion(ay, 5, "1. Qué va en cada archivo")
    for titulo, txt in (("Cuadro de horas", "Es la fuente principal: las clases, sus fechas, módulos, profesores y horas. Si algo de eso cambia, se cambia allá, nunca aquí."),
                        ("Este archivo", "Solo los nombres de los estudiantes y un 1 o un 0 por clase. Es el único lugar donde están los nombres."),
                        ("Sitio de administración", "Junta las dos cosas y muestra faltas permitidas, estados y alertas. Nunca muestra nombres: solo el número de lista, por ejemplo Estudiante 7.")):
        _texto(ay, f, 2, 2, titulo, 40, negrita=True, color=AZUL, fondo=GRIS)
        _texto(ay, f, 3, 5, txt, 40, fondo=GRIS)
        f += 1
    f = _seccion(ay, f + 1, "2. Para empezar, una sola vez")
    f = _paso(ay, f, "1", "En la hoja Asistencia escribe los nombres en la columna Nombre del estudiante. Puedes pegar toda la lista desde el Excel de matriculados. El documento es opcional.")
    f = _paso(ay, f, "2", "En la hoja Resumen escribe en Clases del curso cuántas clases hay en el Cuadro de horas (el número de la última). Sirve para los colores de este archivo; el sitio ya lo toma del Cuadro.")
    f = _seccion(ay, f + 1, "3. Después de cada clase")
    f = _paso(ay, f, "1", "En la hoja Asistencia, en la columna de esa clase, escribe 1 a quien asistió y 0 a quien faltó. No dejes casillas vacías en una clase ya dictada.")
    f = _paso(ay, f, "2", "Guarda. No hay botón de enviar: el sitio de administración se actualiza solo en unos minutos.")
    f = _paso(ay, f, "3", "Mira la columna Estado. A quien aparece en Atención o En riesgo conviene escribirle antes de que pierda. Anota en Seguimiento cómo va el contacto.")
    f = _seccion(ay, f + 1, "4. Qué significa cada estado")
    for texto, fondo, color in (("Al día: ha faltado poco, lejos del límite.", "E2F0D9", "375623"),
                                ("Atención: ya usó la mitad o más de las faltas permitidas.", "FFF2CC", "7F6000"),
                                ("En riesgo: le queda una falta o ninguna. Una más y pierde el certificado.", "FBE0D6", "C00000"),
                                ("Perdió el mínimo: ya no puede llegar al porcentaje mínimo, aunque asista a todo lo que queda.", "C00000", "FFFFFF"),
                                ("Aprobó: terminó el curso y cumplió el mínimo.", "DDEBF7", "1F3864")):
        _texto(ay, f, 2, 5, texto, 26, negrita=True, color=color, fondo=fondo)
        f += 1
    _texto(ay, f, 2, 5, "Los avisos se calculan sobre las faltas permitidas del curso, no sobre un número fijo: sirven igual para 35 clases que para 8. Para certificar se necesita asistir al 80 % de las clases.", 40, color="595959")
    f += 1
    f = _seccion(ay, f + 1, "5. Cómo se leen las columnas grises (no se escriben)")
    for txt in ("Faltas que le quedan: cuántas veces más puede faltar sin perder el mínimo.",
                "Puede llegar hasta: el mejor porcentaje posible si no falta a nada más.",
                "% hasta hoy: su asistencia solo con las clases ya dictadas.",
                "Sin marcar: clases ya dictadas con su casilla vacía. Debe ser 0: si no, falta marcarla.",
                "Seguimiento: nota libre (Contactado, En conversación, Resuelto) para el equipo."):
        _texto(ay, f, 2, 5, "•  " + txt, 24)
        f += 1
    f = _seccion(ay, f + 1, "6. Cuando algo cambia")
    for txt in ("Se agrega una clase o cambia una fecha: se hace en el Cuadro de horas. El sitio la toma sola. Aquí solo sube Clases del curso en la hoja Resumen, si la clase es nueva.",
                "Una clase se reprograma: cambia su fecha en el Cuadro; la clase conserva su número y sus marcas. Deja sus casillas vacías hasta que se dicte.",
                "Una clase se cancela: no la borres del Cuadro por tu cuenta, porque las demás cambiarían de número y las marcas quedarían desplazadas. Avisa a quien administra el sitio.",
                "Una clase pasa a Dictada sola en cuanto tiene alguna marca."):
        _texto(ay, f, 2, 5, "•  " + txt, 40)
        f += 1
    f = _seccion(ay, f + 1, "7. Estudiantes nuevos o que se retiran")
    for txt in ("Para agregar a alguien, escribe su nombre en la primera fila libre debajo de la lista. La tabla se extiende y las cuentas se llenan solas.",
                "Si alguien se retira, no borres su fila: deja sus marcas. Si hay que quitarlo de las cuentas, borra solo su nombre; lo demás queda en el historial de versiones de SharePoint."):
        _texto(ay, f, 2, 5, "•  " + txt, 40)
        f += 1
    f = _seccion(ay, f + 1, "8. Importante")
    for txt in ("No cambies los nombres de las columnas Clase 1, Clase 2, etc. ni los de las columnas grises: el sitio los usa para leer el archivo.",
                "La hoja Cronograma es opcional y solo sirve para ver fechas sobre cada columna. Las fechas oficiales son las del Cuadro de horas.",
                "El archivo cabe hasta 40 clases. Las columnas que sobran aparecen en gris, como Sin programar."):
        _texto(ay, f, 2, 5, "•  " + txt, 40)
        f += 1
    ay.page_setup.orientation, ay.page_setup.fitToWidth, ay.page_setup.fitToHeight = "portrait", 1, 0
    ay.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def construir(d, salida):
    wb = Workbook()
    clases = d.get("clases", [])
    _hoja_asistencia(wb, d["curso"], d.get("estudiantes", []), bool(clases))
    _hoja_resumen(wb, len(clases) or None, d.get("minimo", 0.8))
    _hoja_cronograma(wb, clases)
    _hoja_instrucciones(wb, d["curso"])
    wb.move_sheet("Instrucciones", offset=-3)
    wb.active = 1
    wb.save(salida)


def main(argv):
    if argv and argv[0] == "--plantilla":
        construir({"curso": "Nombre del curso"}, argv[1])
        return
    construir(json.loads(Path(argv[0]).read_text(encoding="utf-8")), argv[1])


if __name__ == "__main__":
    main(sys.argv[1:])
