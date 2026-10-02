#!/usr/bin/env python3
"""Construye data/data.json, estado/cambios.json y cronograma.ics de UN curso desde el Excel (o JSON crudo del Plan B).

Uso: python scripts/construir_datos.py ENTRADA [--json] [--raiz c/<curso>] [--modificado ISO]
--raiz es la carpeta del curso (con config/contenido.json); por defecto la raíz del repositorio (curso único).
Sale con código 1 y mensaje claro ante un error bloqueante (se conserva el último data.json bueno).
Con hash sin cambios no escribe nada (código 0, "SIN CAMBIOS").
"""
import argparse, hashlib, json, re, sys, unicodedata
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

HOJA = "DISTRIBUCIÓN HORAS"
COLS = ["No.", "Unidad", "Tema", "Horas", "No. de clase", "Profesor", "Día", "Fecha", "Asistente PAT", "Link de reunión"]
DIAS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
BOGOTA = timezone(timedelta(hours=-5))  # sin horario de verano


class ErrorDatos(Exception):
    """Error que detiene la publicación."""


def slug(s):
    s = unicodedata.normalize("NFD", s.lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")[:60]


def limpiar(s):
    return re.sub(r"[ \t ]+", " ", str(s)).strip() if s is not None else ""


def puntos(tema):
    partes = [p.strip() for p in re.split(r"\s{2,}|\n", str(tema)) if p.strip()]
    out, i = [], 0
    while i < len(partes):
        p = partes[i].lstrip("-–— ").strip()
        if p.upper() == "CASO" and i + 1 < len(partes):
            out.append("CASO: " + partes[i + 1].lstrip("-–— ").strip())
            i += 2
            continue
        if p:
            out.append(p)
        i += 1
    return out


def fecha_a_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return (datetime(1899, 12, 30) + timedelta(days=int(v))).date()
    if isinstance(v, str):
        if re.fullmatch(r"\d+(\.\d+)?", v.strip()):  # serial de Excel como texto (así lo entrega Power Automate)
            return fecha_a_date(float(v))
        try:
            return date.fromisoformat(v.strip()[:10])
        except ValueError:
            pass
    raise ErrorDatos(f"Fecha no interpretable: {v!r}")


def leer_xlsx(ruta):
    from openpyxl import load_workbook
    if Path(ruta).read_bytes()[:2] != b"PK":
        raise ErrorDatos("El archivo descargado no es un .xlsx (suele ser un vínculo sin permiso o que exige iniciar sesión).")
    try:
        wb = load_workbook(ruta, data_only=True)
    except Exception as e:
        raise ErrorDatos(f"No se pudo abrir el .xlsx: {e}")
    if HOJA not in wb.sheetnames:
        raise ErrorDatos(f"No existe la hoja '{HOJA}'. Hojas encontradas: {wb.sheetnames}")
    ws = wb[HOJA]
    hdr = next((r for r in range(1, 30) if [limpiar(c.value) for c in ws[r][:10]] == COLS), None)
    if hdr is None:
        vistos = {limpiar(c.value) for r in range(1, 15) for c in ws[r][:10]}
        faltan = [c for c in COLS if c not in vistos]
        raise ErrorDatos(f"No se encontró la fila de encabezados exactos A:J. Columnas que faltan: {faltan or 'orden distinto'}")
    # solo A:J, se ignora todo lo que esté fuera
    return [[ws.cell(r, c).value for c in range(1, 11)] for r in range(hdr + 1, ws.max_row + 1)]


def leer_json(ruta):
    """Plan B: lista de objetos con los encabezados del Excel como llaves."""
    datos = json.loads(Path(ruta).read_text(encoding="utf-8"))
    if isinstance(datos, dict):
        datos = datos.get("filas") or datos.get("value") or []
    # Excel Online codifica los caracteres especiales de los encabezados: "No." llega como "No_x002e_"
    deco = lambda k: re.sub(r"_x([0-9a-fA-F]{4})_", lambda m: chr(int(m.group(1), 16)), k)
    datos = [{deco(k): v for k, v in d.items()} for d in datos]
    if datos and not set(COLS) <= set(datos[0]):
        raise ErrorDatos(f"El JSON no trae las columnas obligatorias: {[c for c in COLS if c not in datos[0]]}")
    return [[d.get(c) for c in COLS] for d in datos]


def hhmm_a_min(h):
    a, b = h.split(":")
    return int(a) * 60 + int(b)


def min_a_hhmm(m):
    return f"{m // 60:02d}:{m % 60:02d}"


def asignar_horarios(clases, cfg):
    """Pone inicio y fin a cada clase (hora de Colombia).

    modo "ventana" (por defecto): todas las clases usan horario.inicio a horario.fin.
    modo "porHoras": el inicio depende del día (horario.porDia, si no horario.inicio) y el fin es inicio más las horas de la clase
    (horario.duracionPorDefecto si la casilla está vacía). Las clases del mismo día van una tras otra.
    """
    h, des = cfg["horario"], cfg["horario"]["desfase"]
    sig = {}  # fecha -> minuto en que termina la clase anterior del mismo día
    for c in clases:
        if h.get("modo") == "porHoras":
            ini = sig.get(c["fecha"], hhmm_a_min(h.get("porDia", {}).get(c["diaCalculado"], h["inicio"])))
            fin = ini + int(c["horas"] or h.get("duracionPorDefecto", 3)) * 60
            sig[c["fecha"]] = fin
            fin = min(fin, 24 * 60 - 1)  # nunca pasa de medianoche
            i, f = min_a_hhmm(ini), min_a_hhmm(fin)
        else:
            i, f = h["inicio"], h["fin"]
        c["inicio"], c["fin"] = f"{c['fecha']}T{i}:00{des}", f"{c['fecha']}T{f}:00{des}"


def entero_o_none(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        n = float(v)
    except (TypeError, ValueError):
        raise ErrorDatos(f"'Horas' no es un número: {v!r}")
    return int(n) if n == int(n) and n > 0 else None


def filas_a_clases(filas, cfg):
    clases, avisos, ant_tema, ant_raw = [], [], "", ""
    for v in filas:
        if not v[4] or not str(v[4]).strip().upper().startswith("CLASE"):
            continue  # fila TOTAL, vacías y residuos
        try:
            num = int(float(v[0]))
        except (TypeError, ValueError):
            raise ErrorDatos(f"{limpiar(v[4])}: 'No.' no es un número ({v[0]!r})")
        horas = entero_o_none(v[3])
        tema, heredado = limpiar(v[2]), False
        if not tema:
            tema, heredado = ant_tema, True
        else:
            ant_tema, ant_raw = tema, str(v[2])  # el raw conserva los espacios dobles que separan puntos
        d = fecha_a_date(v[7])
        calc = DIAS[d.weekday()]
        c = {
            "id": num, "clase": limpiar(v[4]), "unidad": limpiar(v[1]), "unidadSlug": slug(limpiar(v[1])),
            "tema": tema, "temaPuntos": puntos(ant_raw), "temaHeredado": heredado, "horas": horas,
            "profesor": limpiar(v[5]),
            "profesores": [p for p in (limpiar(x) for x in re.split(r"\s+y\s+", limpiar(v[5]))) if p],
            "dia": limpiar(v[6]), "diaCalculado": calc, "fecha": d.isoformat(),
            "asistentePat": limpiar(v[8]), "linkTeams": limpiar(v[9]), "cambio": None,
        }
        if c["dia"] and c["dia"] != calc:
            avisos.append(f"{c['clase']} dice '{c['dia']}' pero {c['fecha']} es {calc} (se muestra {calc})")
        if c["linkTeams"] and not c["linkTeams"].startswith("https://"):
            avisos.append(f"{c['clase']} tiene un enlace de Teams que no empieza por https:// (se ignora)")
            c["linkTeams"] = ""
        clases.append(c)
    if not clases:
        raise ErrorDatos("No se encontró ninguna clase (filas con 'CLASE' en la columna 'No. de clase').")
    ids = [c["id"] for c in clases]
    dup = sorted({i for i in ids if ids.count(i) > 1})
    if dup:
        raise ErrorDatos(f"Identificadores 'No.' duplicados: {dup}")
    asignar_horarios(clases, cfg)
    # Resumen de lo que falta (las casillas vacías se muestran como "Por confirmar" en la página)
    for etiqueta, vacio in (("profesor", lambda c: not c["profesores"]), ("asistente PAT", lambda c: not c["asistentePat"]),
                            ("enlace de Teams", lambda c: not c["linkTeams"]), ("horas", lambda c: c["horas"] is None)):
        n = sum(1 for c in clases if vacio(c))
        if n:
            avisos.append(f"{n} de {len(clases)} clases sin {etiqueta}")
    return clases, avisos


def hoy_bogota(ahora):
    return ahora.astimezone(BOGOTA).date()


def aplicar_cambios(clases, estado, ahora, dias_aviso):
    """Compara fechas contra el estado previo. Devuelve (estado nuevo, cambios activos)."""
    iso = ahora.strftime("%Y-%m-%dT%H:%M:%SZ")
    actual = {str(c["id"]): c["fecha"] for c in clases}
    if not estado:  # primera ejecución: línea base sin avisos
        estado = {"base": dict(actual), "original": dict(actual), "secuencia": {k: 0 for k in actual}, "cambios": []}
    else:
        estado = json.loads(json.dumps(estado))
        for k, f in actual.items():
            if k not in estado["base"]:  # clase agregada
                estado["base"][k] = estado["original"][k] = f
                estado["secuencia"][k] = 0
                continue
            prev = estado["base"][k]
            if f == prev:
                continue
            activo = next((x for x in estado["cambios"] if str(x["id"]) == k and not x.get("resuelto")), None)
            estado["secuencia"][k] = estado["secuencia"].get(k, 0) + 1
            if f == estado["original"][k]:
                if activo:
                    activo["resuelto"] = True
                    activo["resueltoEn"] = iso
            elif activo:  # segundo cambio seguido: se conserva la fecha original
                activo.update(fechaAnterior=prev, fechaNueva=f, detectadoEn=iso)
            else:
                cl = next(c for c in clases if str(c["id"]) == k)
                estado["cambios"].append({"id": int(k), "clase": cl["clase"], "fechaOriginal": estado["original"][k],
                                          "fechaAnterior": prev, "fechaNueva": f, "detectadoEn": iso, "resuelto": False})
            estado["base"][k] = f
        for k in [k for k in estado["base"] if k not in actual]:  # clase eliminada
            for d in (estado["base"], estado["original"], estado["secuencia"]):
                d.pop(k, None)
            estado["cambios"] = [x for x in estado["cambios"] if str(x["id"]) != k]
    hoy = hoy_bogota(ahora)
    limite = ahora - timedelta(days=dias_aviso)
    activos = []
    for x in estado["cambios"]:
        if x.get("resuelto"):
            continue
        det = datetime.strptime(x["detectadoEn"], "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=timezone.utc)
        if det < limite or date.fromisoformat(x["fechaNueva"]) < hoy:
            continue  # vencido por tiempo o la clase ya pasó
        activos.append(x)
    por_id = {x["id"]: x for x in activos}
    for c in clases:
        x = por_id.get(c["id"])
        if x:
            c["cambio"] = {"tipo": "fecha", "fechaAnterior": x["fechaAnterior"], "fechaOriginal": x["fechaOriginal"], "detectadoEn": x["detectadoEn"]}
        c["secuencia"] = estado["secuencia"].get(str(c["id"]), 0)
    return estado, activos


def armar(clases, cfg, cambios, ahora, modificado=""):
    unidades, vistos = [], {}
    for c in clases:
        if not c["unidadSlug"]:
            continue  # sin unidad en el Excel: la página oculta el temario por unidades
        u = vistos.get(c["unidadSlug"])
        if not u:
            u = vistos[c["unidadSlug"]] = {"slug": c["unidadSlug"], "nombre": c["unidad"], "clases": [], "horas": 0}
            unidades.append(u)
        u["clases"].append(c["id"])
        u["horas"] += c["horas"] or 0
    base = {"clases": clases, "programa": cfg["programa"], "horario": cfg["horario"]}
    hash_ = hashlib.sha256(json.dumps(base, sort_keys=True, ensure_ascii=False).encode()).hexdigest()[:16]
    return {
        "meta": {
            "programa": cfg["programa"], "tipo": cfg.get("tipo", "Curso"), "facultad": cfg["facultad"], "modalidad": cfg["modalidad"],
            "zona": cfg["horario"]["zona"], "totalClases": len(clases),
            "horasTotales": sum(c["horas"] or 0 for c in clases), "horasCompletas": all(c["horas"] for c in clases),
            "inicio": min(c["fecha"] for c in clases), "fin": max(c["fecha"] for c in clases),
            "hash": hash_, "generadoEn": ahora.strftime("%Y-%m-%dT%H:%M:%SZ"), "fuenteModificadaEn": modificado,
        },
        "clases": clases, "unidades": unidades,
        "profesores": sorted({p for c in clases for p in c["profesores"]}),
        "asistentesPat": sorted({c["asistentePat"] for c in clases if c["asistentePat"]}),
        "cambios": cambios,
    }


def construir(entrada, raiz, modo_json=False, ahora=None, modificado=""):
    """Devuelve (datos | None si sin cambios, avisos)."""
    raiz = Path(raiz)
    ahora = ahora or datetime.now(timezone.utc)
    cfg = json.loads((raiz / "config/contenido.json").read_text(encoding="utf-8"))
    filas = leer_json(entrada) if modo_json else leer_xlsx(entrada)
    clases, avisos = filas_a_clases(filas, cfg)
    ruta_estado = raiz / "estado/cambios.json"
    estado = json.loads(ruta_estado.read_text(encoding="utf-8")) if ruta_estado.exists() else None
    estado, activos = aplicar_cambios(clases, estado, ahora, cfg["diasAvisoCambio"])
    datos = armar(clases, cfg, activos, ahora, modificado)
    ruta_datos = raiz / "data/data.json"
    if ruta_datos.exists() and json.loads(ruta_datos.read_text(encoding="utf-8"))["meta"]["hash"] == datos["meta"]["hash"]:
        return None, avisos
    from generar_ics import generar
    for d in (ruta_datos.parent, ruta_estado.parent):
        d.mkdir(exist_ok=True)
    ruta_estado.write_text(json.dumps(estado, ensure_ascii=False, indent=2), encoding="utf-8")
    ruta_datos.write_text(json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8")
    (raiz / "cronograma.ics").write_bytes(generar(datos, cfg).encode("utf-8"))
    return datos, avisos


def main():
    sys.path.insert(0, str(Path(__file__).parent))
    ap = argparse.ArgumentParser()
    ap.add_argument("entrada")
    ap.add_argument("--json", action="store_true", help="Plan B: la entrada es JSON crudo, no .xlsx")
    ap.add_argument("--raiz", default=str(Path(__file__).resolve().parent.parent))
    ap.add_argument("--modificado", default="")
    a = ap.parse_args()
    ruta_error = Path(a.raiz) / "estado/error.json"
    try:
        datos, avisos = construir(a.entrada, a.raiz, a.json, modificado=a.modificado)
    except ErrorDatos as e:
        print(f"ERROR: {e}", file=sys.stderr)
        ruta_error.parent.mkdir(exist_ok=True)  # el sitio base de administración lo muestra; se conserva el último data.json bueno
        ruta_error.write_text(json.dumps({"mensaje": str(e), "en": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")}, ensure_ascii=False), encoding="utf-8")
        sys.exit(1)
    ruta_error.unlink(missing_ok=True)
    for av in avisos:
        print(f"AVISO: {av}")
    if datos is None:
        print("SIN CAMBIOS: el hash coincide, no se publica")
    else:
        m = datos["meta"]
        for x in datos["cambios"]:
            print(f"CAMBIO: {x['clase']} pasó de {x['fechaAnterior']} a {x['fechaNueva']} (original {x['fechaOriginal']})")
        print(f"OK: {m['totalClases']} clases, {m['horasTotales']} horas{'' if m['horasCompletas'] else ' (incompletas)'}, {len(datos['unidades'])} unidades, "
              f"{m['inicio']} a {m['fin']}, {len(datos['cambios'])} cambios activos")


if __name__ == "__main__":
    main()
