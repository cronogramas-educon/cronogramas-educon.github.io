#!/usr/bin/env python3
"""Genera el Excel "Registro de cursos" a partir de config/cursos.json: python scripts/generar_registro_xlsx.py SALIDA.xlsx"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parent))
import cursos as cu

COLS = [("Periodo", 11), ("Carpeta del periodo", 24), ("Carpeta del curso", 44), ("Tipo", 13), ("Nombre del curso", 44), ("Nombre completo", 70),
        ("Modalidad", 26), ("Hora de inicio", 14), ("Hora de fin", 13), ("Inicio del sábado", 17), ("Página oficial", 60)]

AYUDA = [
    "REGISTRO DE CURSOS Y DIPLOMADOS",
    "",
    "Este archivo decide qué cursos tiene el sitio de cronogramas. Cada fila es un curso de un periodo. No hace falta tocar GitHub ni pedir ayuda: al guardar, el sitio se actualiza solo en un par de minutos.",
    "",
    "CÓMO ABRIR UN PERIODO NUEVO",
    "1. En SharePoint crea la carpeta del periodo (por ejemplo EDU CONTINUA 2027-1) y, adentro, la carpeta de cada curso con su Excel Cuadro de horas módulos y profesores.",
    "2. En la hoja Registro copia las filas del periodo anterior y pégalas debajo, dentro de la tabla.",
    "3. En las filas nuevas cambia el Periodo (por ejemplo 2027-1) y la Carpeta del periodo. Cambia también lo que sea distinto en cada curso.",
    "4. Guarda. En el sitio de administración aparece cada curso nuevo con su enlace para estudiantes.",
    "5. Guarda también el Excel de cada curso para que se publique su cronograma.",
    "",
    "CÓMO AGREGAR UN CURSO NUEVO O CAMBIAR UNO",
    "Escribe una fila nueva con todos los datos. Para cambiar un curso, edita su fila. El curso se reconoce por Carpeta del periodo más Carpeta del curso, así que esas dos columnas no se deben cambiar en un curso que ya existe.",
    "",
    "QUÉ ESCRIBIR EN CADA COLUMNA",
    "Periodo: el semestre o ciclo, por ejemplo 2026-2.",
    "Carpeta del periodo y Carpeta del curso: los nombres exactos de las carpetas en SharePoint, tal como se ven. Si el curso está dentro de otra carpeta, escribe las dos separadas por una barra, por ejemplo IA en el seclegal III Cohorte/GRUPO 1- Laura.",
    "Tipo: Diplomado o Curso.",
    "Nombre del curso: el nombre corto que verán los estudiantes como título.",
    "Nombre completo: el nombre oficial del programa. Si lo dejas vacío se usa el nombre del curso.",
    "Modalidad: por ejemplo Remota (Microsoft Teams). Si lo dejas vacío aparece Por confirmar.",
    "Hora de inicio y Hora de fin: por ejemplo 18:00 y 21:00. Si dejas vacía la Hora de fin, cada clase termina según sus horas del Excel del curso.",
    "Inicio del sábado: solo si las clases de sábado empiezan a otra hora y dejaste vacía la Hora de fin, por ejemplo 08:00.",
    "Página oficial: la dirección de la página del programa en unisabana.edu.co (opcional).",
    "",
    "IMPORTANTE",
    "Nunca borres una fila de un curso que ya se dictó: ese curso pasa solo al archivo del sitio y su enlace sigue funcionando en modo lectura. Si borras la fila no pasa nada con la página, pero se pierde el rastro de que existió.",
    "Si una fila tiene un error (falta un dato, una hora mal escrita), el sitio de administración lo avisa en la sección Cursos y periodos y no se procesa esa fila.",
    "Los códigos de las direcciones de los estudiantes los genera el sistema. No se escriben aquí.",
]


def main(salida):
    reg = cu.registro()
    wb = Workbook()
    ws = wb.active
    ws.title = "Registro"
    for j, (t, w) in enumerate(COLS, 1):
        c = ws.cell(1, j, t)
        c.font, c.fill = Font(bold=True, name="Calibri", size=11), PatternFill("solid", fgColor="5B9BD5")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 24
    for i, c in enumerate(reg["cursos"], 2):
        per, cur = c["carpeta"].split("/", 1)
        h = c["horario"]
        fila = [c["periodo"], per, cur, c["tipo"], c["programaCorto"], c["programa"], c["modalidad"], h["inicio"], h.get("fin", ""),
                h.get("porDia", {}).get("Sábado", ""), c.get("urlPaginaOficial", "")]
        for j, v in enumerate(fila, 1):
            x = ws.cell(i, j, v)
            x.alignment = Alignment(vertical="center", wrap_text=True)
            if j in (8, 9, 10):
                x.number_format = "@"
    n = len(reg["cursos"]) + 1
    t = Table(displayName="Table1", ref=f"A1:K{n}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium20", showRowStripes=True)
    ws.add_table(t)
    for col in "HIJ":  # el formato de texto se extiende a las filas que se agreguen
        for r in range(n + 1, 200):
            ws[f"{col}{r}"].number_format = "@"
    dv = DataValidation(type="list", formula1='"Diplomado,Curso"', allow_blank=True)
    dv.error, dv.errorTitle = "Escribe Diplomado o Curso", "Tipo"
    ws.add_data_validation(dv)
    dv.add("D2:D200")
    ws.freeze_panes = "A2"
    ay = wb.create_sheet("Instrucciones")
    ay.column_dimensions["A"].width = 130
    for i, l in enumerate(AYUDA, 1):
        c = ay.cell(i, 1, l)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if l.isupper() and l:
            c.font = Font(bold=True, size=12 if i > 1 else 14)
    wb.move_sheet("Instrucciones", offset=-1)
    wb.active = 1
    wb.save(salida)


if __name__ == "__main__":
    main(sys.argv[1])
