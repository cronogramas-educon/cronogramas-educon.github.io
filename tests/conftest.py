"""Sitio de prueba: dos cursos con datos sanitizados (enlaces de ejemplo) para las pruebas del frontend."""
import json, shutil, sys
from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "scripts"))
import construir_datos as cd
import cursos as cu

FIX = RAIZ / "tests/fixtures/muestra_sanitizada.xlsx"
T0 = datetime(2026, 10, 1, 12, 0, tzinfo=timezone.utc)
A, B = "compliance-anticorrupcion", "derecho-laboral"


def _excel_incompleto(destino):
    """Curso a medio llenar: sin unidad, asistente, enlaces ni docente en las primeras clases, y horario por horas."""
    wb = load_workbook(FIX)
    ws = wb[cd.HOJA]
    for r in range(3, 38):
        for col in (2, 9, 10):
            ws.cell(r, col).value = None
    for r in (3, 4, 5):
        ws.cell(r, 6).value = None
    ws["D3"].value = None                       # clase 1 sin horas
    ws["D5"].value = 4                          # clase 3 (sábado) de 4 horas
    ws["H4"].value = ws["H3"].value             # clase 2 el mismo jueves que la 1
    ws["H5"].value = datetime(2026, 10, 3)      # clase 3 en sábado
    wb.save(destino)


@pytest.fixture(scope="session")
def sitio(tmp_path_factory):
    t = tmp_path_factory.mktemp("repo")
    reg = cu.registro()
    reg["cursos"] = [c for c in reg["cursos"] if c["id"] in (A, B)]
    for d in ("css", "js", "assets", "plantilla"):
        shutil.copytree(RAIZ / d, t / d)
    (t / "config").mkdir()
    (t / "config/cursos.json").write_text(json.dumps(reg, ensure_ascii=False), encoding="utf-8")
    cu.sincronizar(t)
    cd.construir(FIX, t / "cursos" / A, ahora=T0)
    _excel_incompleto(t / "incompleto.xlsx")
    cd.construir(t / "incompleto.xlsx", t / "cursos" / B, ahora=T0)
    (t / "cursos" / B / "estado/error.json").write_text(json.dumps({"mensaje": "Falta la columna Profesor", "en": "2026-10-01T12:00:00Z"}), encoding="utf-8")
    out = cu.armar_sitio(t / "_sitio", "test", raiz=t)
    por = {c["id"]: c for c in reg["cursos"]}
    return {"dir": out, "a": cu.sitio_ruta(por[A]), "b": cu.sitio_ruta(por[B]), "hub": f"g/{reg['global']['adminCodigo']}",
            "datos_a": t / "cursos" / A / "data/data.json", "cfg_a": t / "cursos" / A / "config/contenido.json", "reg": reg}
