import json, shutil, sys
from datetime import datetime, timezone
from pathlib import Path

import pytest
from icalendar import Calendar
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "scripts"))
import construir_datos as cd
from generar_ics import generar, plegar

FIX = RAIZ / "tests/fixtures/muestra_sanitizada.xlsx"
CURSO = "compliance-anticorrupcion"
CFG_RUTA = RAIZ / "cursos" / CURSO / "config/contenido.json"
CFG = json.loads(CFG_RUTA.read_text(encoding="utf-8"))
T0 = datetime(2026, 10, 1, 12, 0, tzinfo=timezone.utc)


@pytest.fixture
def raiz(tmp_path):
    (tmp_path / "config").mkdir()
    shutil.copy(CFG_RUTA, tmp_path / "config")
    return tmp_path


def xlsx_con(tmp_path, nombre, editar):
    wb = load_workbook(FIX)
    editar(wb[cd.HOJA])
    p = tmp_path / nombre
    wb.save(p)
    return p


def fila_de(ws, n):
    return next(r for r in range(3, 40) if ws.cell(r, 1).value == n)


def test_muestra_cumple_aceptacion(raiz):
    datos, avisos = cd.construir(FIX, raiz, ahora=T0)
    m = datos["meta"]
    assert (m["totalClases"], m["horasTotales"], len(datos["unidades"])) == (35, 105, 17)
    assert (m["inicio"], m["fin"]) == ("2026-10-01", "2026-12-04")
    assert [c["id"] for c in datos["clases"] if c["temaHeredado"]] == [17, 19, 22, 25, 26]
    assert all(c["profesor"] == "SABANA" for c in datos["clases"][1:6])
    assert "abril" not in json.dumps(datos, ensure_ascii=False).lower()
    assert avisos == []
    c23 = datos["clases"][22]
    assert c23["profesores"] == ["Marta Lucia Ramírez", "Beatriz Londoño P"]
    assert c23["inicio"].endswith("T17:00:00-05:00")


def test_fixture_sin_enlaces_reales():
    ws = load_workbook(FIX)[cd.HOJA]
    assert {ws.cell(r, 10).value for r in range(3, 38)} == {"https://teams.microsoft.com/meet/EJEMPLO"}


def test_sin_cambios_no_reescribe(raiz):
    cd.construir(FIX, raiz, ahora=T0)
    datos, _ = cd.construir(FIX, raiz, ahora=T0)
    assert datos is None


def test_puntos_caso():
    assert cd.puntos("Tema A  CASO  Odebrecht\n- otro") == ["Tema A", "CASO: Odebrecht", "otro"]


def test_fecha_serial_e_iso():
    assert cd.fecha_a_date(46296).isoformat() == "2026-10-01"
    assert cd.fecha_a_date("2026-10-01").isoformat() == "2026-10-01"
    assert cd.fecha_a_date("46296").isoformat() == "2026-10-01"
    with pytest.raises(cd.ErrorDatos):
        cd.fecha_a_date("pronto")


def test_modo_json_plan_b(raiz, tmp_path):
    datos0, _ = cd.construir(FIX, raiz, ahora=T0)
    ws = load_workbook(FIX)[cd.HOJA]
    filas = [dict(zip(cd.COLS, [ws.cell(r, c).value for c in range(1, 11)])) for r in range(3, 39)]
    for f in filas:
        if hasattr(f["Fecha"], "date"):
            f["Fecha"] = (f["Fecha"].date() - datetime(1899, 12, 30).date()).days  # serial, como Excel Online
    p = tmp_path / "crudo.json"
    p.write_text(json.dumps(filas, default=str), encoding="utf-8")
    r2 = tmp_path / "r2"
    (r2 / "config").mkdir(parents=True)
    shutil.copy(CFG_RUTA, r2 / "config")
    datos, _ = cd.construir(p, r2, modo_json=True, ahora=T0)
    assert datos["meta"]["hash"] == datos0["meta"]["hash"]


# --- errores bloqueantes
def test_no_es_xlsx(raiz, tmp_path):
    p = tmp_path / "x.xlsx"
    p.write_text("<html>inicia sesión</html>")
    with pytest.raises(cd.ErrorDatos, match="no es un .xlsx"):
        cd.construir(p, raiz)


def test_falta_encabezado(raiz, tmp_path):
    p = xlsx_con(tmp_path, "a.xlsx", lambda ws: setattr(ws["F2"], "value", None))
    with pytest.raises(cd.ErrorDatos, match="Profesor"):
        cd.construir(p, raiz)


def test_cero_clases(raiz, tmp_path):
    def vaciar(ws):
        for r in range(3, 38):
            ws.cell(r, 5).value = None
    with pytest.raises(cd.ErrorDatos, match="ninguna clase"):
        cd.construir(xlsx_con(tmp_path, "b.xlsx", vaciar), raiz)


def test_fecha_invalida_y_duplicado(raiz, tmp_path):
    p = xlsx_con(tmp_path, "c.xlsx", lambda ws: setattr(ws["H3"], "value", "pronto"))
    with pytest.raises(cd.ErrorDatos, match="Fecha"):
        cd.construir(p, raiz)
    p = xlsx_con(tmp_path, "d.xlsx", lambda ws: setattr(ws["A4"], "value", 1))
    with pytest.raises(cd.ErrorDatos, match="duplicados"):
        cd.construir(p, raiz)


def test_error_conserva_ultimo_bueno(raiz, tmp_path):
    cd.construir(FIX, raiz, ahora=T0)
    antes = (raiz / "data/data.json").read_text(encoding="utf-8")
    p = xlsx_con(tmp_path, "e.xlsx", lambda ws: setattr(ws["F2"], "value", None))
    with pytest.raises(cd.ErrorDatos):
        cd.construir(p, raiz)
    assert (raiz / "data/data.json").read_text(encoding="utf-8") == antes


# --- advertencias
def test_advertencias(raiz, tmp_path):
    def romper(ws):
        ws["G3"] = "Lunes"      # día no coincide
        ws["J4"] = None         # sin link
    datos, avisos = cd.construir(xlsx_con(tmp_path, "f.xlsx", romper), raiz, ahora=T0)
    assert len(avisos) == 2 and datos["clases"][0]["diaCalculado"] == "Jueves"
    assert any("Lunes" in a for a in avisos) and "1 de 35 clases sin enlace de Teams" in avisos


def test_casillas_vacias_no_rompen(raiz, tmp_path):
    """Un Excel a medio llenar es normal: profesor, asistente, enlace, horas y unidad en blanco se publican como 'por confirmar'."""
    def vaciar(ws):
        for r in range(3, 38):
            for col in (2, 4, 6, 9, 10):  # Unidad, Horas, Profesor, Asistente PAT, Link
                ws.cell(r, col).value = None
    datos, avisos = cd.construir(xlsx_con(tmp_path, "v.xlsx", vaciar), raiz, ahora=T0)
    assert datos["unidades"] == [] and datos["profesores"] == [] and datos["asistentesPat"] == []
    assert all(c["horas"] is None and c["profesores"] == [] and c["linkTeams"] == "" for c in datos["clases"])
    assert datos["meta"]["horasCompletas"] is False and datos["meta"]["horasTotales"] == 0
    for t in ("sin profesor", "sin asistente PAT", "sin enlace de Teams", "sin horas"):
        assert f"35 de 35 clases {t}" in avisos


def test_enlace_sin_https_se_ignora(raiz, tmp_path):
    datos, avisos = cd.construir(xlsx_con(tmp_path, "w.xlsx", lambda ws: setattr(ws["J3"], "value", "teams.microsoft.com/x")), raiz, ahora=T0)
    assert datos["clases"][0]["linkTeams"] == "" and any("https://" in a for a in avisos)


def test_horario_por_horas_con_sabados_y_clases_el_mismo_dia(tmp_path):
    cfg = json.loads(CFG_RUTA.read_text(encoding="utf-8"))
    cfg["horario"] = {**cfg["horario"], "modo": "porHoras", "inicio": "18:00", "porDia": {"Sábado": "08:00"}, "duracionPorDefecto": 1}
    r = tmp_path / "laboral"
    (r / "config").mkdir(parents=True)
    (r / "config/contenido.json").write_text(json.dumps(cfg), encoding="utf-8")

    def ed(ws):
        ws["H4"].value = ws["H3"].value          # clase 2 el mismo jueves que la 1
        ws["H5"].value = datetime(2026, 10, 3)   # clase 3 en sábado
        ws["D3"].value = None; ws["D4"].value = 3; ws["D5"].value = 4
    datos, _ = cd.construir(xlsx_con(tmp_path, "x.xlsx", ed), r, ahora=T0)
    h = [(c["inicio"][11:16], c["fin"][11:16]) for c in datos["clases"][:3]]
    assert h == [("18:00", "19:00"), ("19:00", "22:00"), ("08:00", "12:00")]


def test_error_se_registra_para_el_sitio_base(raiz, tmp_path, monkeypatch, capsys):
    p = xlsx_con(tmp_path, "z.xlsx", lambda ws: setattr(ws["F2"], "value", None))
    monkeypatch.setattr(sys, "argv", ["construir_datos.py", str(p), "--raiz", str(raiz)])
    with pytest.raises(SystemExit):
        cd.main()
    assert "Profesor" in json.loads((raiz / "estado/error.json").read_text(encoding="utf-8"))["mensaje"]
    monkeypatch.setattr(sys, "argv", ["construir_datos.py", str(FIX), "--raiz", str(raiz)])
    cd.main()
    assert not (raiz / "estado/error.json").exists()


# --- detección de cambios
def cambiar(tmp_path, raiz, nombre, ahora, cambios):
    """cambios: {id: nueva fecha ISO}"""
    def ed(ws):
        for i, f in cambios.items():
            ws.cell(fila_de(ws, i), 8).value = datetime.fromisoformat(f)
    return cd.construir(xlsx_con(tmp_path, nombre, ed), raiz, ahora=ahora)[0]


def dt(d, h=12):
    return datetime(2026, 10, d, h, tzinfo=timezone.utc)


def test_cambio_de_fecha_y_resolucion(raiz, tmp_path):
    cd.construir(FIX, raiz, ahora=T0)  # línea base, sin avisos
    assert json.loads((raiz / "data/data.json").read_text(encoding="utf-8"))["cambios"] == []
    orig = load_workbook(FIX)[cd.HOJA]
    f7 = orig.cell(fila_de(orig, 7), 8).value.date().isoformat()
    d = cambiar(tmp_path, raiz, "g.xlsx", dt(2), {7: "2026-10-14"})
    c7 = d["clases"][6]
    assert c7["cambio"]["fechaAnterior"] == f7 and c7["secuencia"] == 1
    assert d["cambios"][0]["fechaNueva"] == "2026-10-14"
    # vuelve a la original: se resuelve
    d = cambiar(tmp_path, raiz, "h.xlsx", dt(3), {})
    assert d["cambios"] == [] and d["clases"][6]["cambio"] is None and d["clases"][6]["secuencia"] == 2


def test_dos_cambios_seguidos_conservan_original(raiz, tmp_path):
    cd.construir(FIX, raiz, ahora=T0)
    cambiar(tmp_path, raiz, "i.xlsx", dt(2), {7: "2026-10-14"})
    d = cambiar(tmp_path, raiz, "j.xlsx", dt(3), {7: "2026-10-20"})
    x = d["cambios"][0]
    assert (x["fechaAnterior"], x["fechaNueva"]) == ("2026-10-14", "2026-10-20") and x["fechaOriginal"] != "2026-10-14"
    assert len(d["cambios"]) == 1 and d["clases"][6]["secuencia"] == 2


def test_aviso_caduca_a_los_14_dias(raiz, tmp_path):
    cd.construir(FIX, raiz, ahora=T0)
    cambiar(tmp_path, raiz, "k.xlsx", dt(2), {7: "2026-12-01"})
    d = cambiar(tmp_path, raiz, "l.xlsx", datetime(2026, 10, 17, 12, tzinfo=timezone.utc), {7: "2026-12-01"})
    assert d["cambios"] == [] and d["clases"][6]["cambio"] is None


def test_cambio_a_fecha_pasada_no_avisa(raiz, tmp_path):
    cd.construir(FIX, raiz, ahora=datetime(2026, 10, 20, tzinfo=timezone.utc))
    d = cambiar(tmp_path, raiz, "m.xlsx", datetime(2026, 10, 21, tzinfo=timezone.utc), {20: "2026-10-05"})
    assert d["cambios"] == []


def test_clase_agregada_y_eliminada(raiz, tmp_path):
    def sin_35(ws):
        ws.delete_rows(fila_de(ws, 35))
    cd.construir(xlsx_con(tmp_path, "n.xlsx", sin_35), raiz, ahora=T0)  # base de 34
    d, _ = cd.construir(FIX, raiz, ahora=dt(2))                          # reaparece la 35
    assert d["meta"]["totalClases"] == 35 and d["cambios"] == []
    d, _ = cd.construir(xlsx_con(tmp_path, "o.xlsx", sin_35), raiz, ahora=dt(3))
    estado = json.loads((raiz / "estado/cambios.json").read_text(encoding="utf-8"))
    assert d["meta"]["totalClases"] == 34 and "35" not in estado["base"]


# --- .ics
def test_ics_valido_y_horas(raiz):
    datos, _ = cd.construir(FIX, raiz, ahora=T0)
    ics = (raiz / "cronograma.ics").read_bytes()
    assert b"\r\n" in ics and b"\n" not in ics.replace(b"\r\n", b"")
    for linea in ics.split(b"\r\n"):
        assert len(linea) <= 75
    cal = Calendar.from_ical(ics)
    evs = [e for e in cal.walk("VEVENT")]
    assert len(evs) == 35
    e = evs[0]
    assert str(e["UID"]) == "clase-01@cronograma-compliance-anticorrupcion"
    ini = e["DTSTART"].dt.astimezone(timezone.utc)
    fin = e["DTEND"].dt.astimezone(timezone.utc)
    assert (ini.hour, fin.hour) == (22, 1) and fin.day == 2  # 17:00 a 20:00 Bogotá = 22:00 a 01:00 UTC
    assert str(e["LOCATION"]) == "Microsoft Teams" and "Marta Lucia" in str(e["DESCRIPTION"])
    assert e.walk("VALARM")[0]["TRIGGER"].dt.total_seconds() == -1800


def test_ics_sequence_sube_con_cambio(raiz, tmp_path):
    cd.construir(FIX, raiz, ahora=T0)
    cambiar(tmp_path, raiz, "p.xlsx", dt(2), {7: "2026-10-14"})
    e = [e for e in Calendar.from_ical((raiz / "cronograma.ics").read_bytes()).walk("VEVENT") if "clase-07" in str(e["UID"])][0]
    assert int(e["SEQUENCE"]) == 1 and e["DTSTART"].dt.day == 14


def test_ics_escape_y_plegado_utf8():
    larga = "DESCRIPTION:" + "áéíóú, ñ; " * 30
    plegada = plegar(larga)
    assert all(len(p.encode()) <= 75 for p in plegada.split("\r\n"))
    assert plegada.replace("\r\n ", "") == larga


def test_json_con_encabezados_codificados_de_excel_online(raiz, tmp_path):
    """El conector de Excel Online devuelve 'No.' como 'No_x002e_'."""
    ws = load_workbook(FIX)[cd.HOJA]
    cod = lambda k: k.replace(".", "_x002e_")
    filas = [{cod(c): ws.cell(r, i + 1).value for i, c in enumerate(cd.COLS)} for r in range(3, 38)]
    for f in filas:
        f[cod("Fecha")] = (f[cod("Fecha")].date() - datetime(1899, 12, 30).date()).days
    p = tmp_path / "codificado.json"
    p.write_text(json.dumps(filas, default=str), encoding="utf-8")
    datos, _ = cd.construir(p, raiz, modo_json=True, ahora=T0)
    assert datos["meta"]["totalClases"] == 35 and datos["clases"][0]["id"] == 1


def test_plantilla_maestra_cumple_el_formato(tmp_path):
    import generar_plantilla_xlsx as gp
    from openpyxl import load_workbook
    f = tmp_path / "p.xlsx"
    gp.main(f)
    ws = load_workbook(f)[cd.HOJA]
    assert [ws.cell(2, j).value for j in range(1, 11)] == cd.COLS and "Table1" in ws.tables
    assert len(ws.data_validations.dataValidation) == 5
