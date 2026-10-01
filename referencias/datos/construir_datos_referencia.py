#!/usr/bin/env python3
"""
Parser de REFERENCIA (validado contra la muestra). Claude Code puede reutilizarlo o reescribirlo.
Uso: python construir_datos_referencia.py cronograma_muestra.xlsx data.ejemplo.json
Reglas: ver el prompt (sección "Reglas de limpieza").
"""
import sys, re, json, hashlib, unicodedata
from datetime import datetime, date, timedelta, timezone
from openpyxl import load_workbook

HOJA = "DISTRIBUCIÓN HORAS"
COLS = ["No.", "Unidad", "Tema", "Horas", "No. de clase", "Profesor", "Día", "Fecha", "Asistente PAT", "Link de reunión"]
DIAS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
TZ = "-05:00"  # America/Bogota, sin horario de verano

def slug(s):
    s = unicodedata.normalize("NFD", s.lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")[:60]

def limpiar(s):
    return re.sub(r"[ \t\u00a0]+", " ", str(s)).strip() if s is not None else ""

def puntos(tema):
    partes = [p.strip() for p in re.split(r"\s{2,}|\n", str(tema)) if p.strip()]
    out, i = [], 0
    while i < len(partes):
        p = partes[i].lstrip("-– ").strip()
        if p.upper() == "CASO" and i + 1 < len(partes):
            out.append("CASO: " + partes[i + 1].lstrip("-– ").strip()); i += 2; continue
        out.append(p); i += 1
    return out

def fecha_a_date(v):
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    if isinstance(v, (int, float)): return (datetime(1899, 12, 30) + timedelta(days=int(v))).date()
    raise ValueError(f"Fecha no interpretable: {v!r}")

def main(xlsx, salida):
    ws = load_workbook(xlsx, data_only=True)[HOJA]
    # 1) localizar encabezado por contenido (no por número de fila)
    hdr = next(r for r in range(1, 15) if [limpiar(c.value) for c in ws[r][:10]] == COLS)
    filas, ant_tema, ant_raw = [], "", ""
    for r in range(hdr + 1, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, 11)]  # solo A:J (ignora Q y demás)
        if not v[4] or not str(v[4]).upper().startswith("CLASE"):  # fin de datos / fila TOTAL
            continue
        tema = limpiar(v[2]) if v[2] else ""
        heredado = False
        if not tema: tema, heredado = ant_tema, True
        else: ant_tema, ant_raw = tema, str(v[2])  # ant_raw conserva los espacios dobles para separar puntos
        d = fecha_a_date(v[7])
        dia_calc = DIAS[d.weekday()]
        filas.append({
            "id": int(v[0]), "clase": limpiar(v[4]), "unidad": limpiar(v[1]), "unidadSlug": slug(limpiar(v[1])),
            "tema": tema, "temaPuntos": puntos(ant_raw),
            "temaHeredado": heredado, "horas": int(v[3]), "profesor": limpiar(v[5]),
            "profesores": [limpiar(p) for p in re.split(r"\s+y\s+", limpiar(v[5]))],
            "dia": limpiar(v[6]), "diaCalculado": dia_calc, "fecha": d.isoformat(),
            "inicio": f"{d.isoformat()}T17:00:00{TZ}", "fin": f"{d.isoformat()}T20:00:00{TZ}",
            "asistentePat": limpiar(v[8]), "linkTeams": limpiar(v[9]), "cambio": None,
        })
    unidades, vistos = [], {}
    for f in filas:
        u = vistos.setdefault(f["unidadSlug"], {"slug": f["unidadSlug"], "nombre": f["unidad"], "clases": [], "horas": 0})
        if u not in unidades: unidades.append(u)
        u["clases"].append(f["id"]); u["horas"] += f["horas"]
    datos = {
        "meta": {
            "programa": "Diplomado Compliance Anti-corrupción y Anti-lavado, más allá de las Normas y la Teoría",
            "cohorte": "3er 2026-2", "facultad": "Facultad de Estudios Jurídicos, Políticos e Internacionales",
            "modalidad": "Remota (Microsoft Teams)", "horario": {"inicio": "17:00", "fin": "20:00", "zona": "America/Bogota"},
            "totalClases": len(filas), "horasTotales": sum(f["horas"] for f in filas),
            "inicio": filas[0]["fecha"], "fin": max(f["fecha"] for f in filas),
        },
        "clases": filas, "unidades": unidades,
        "profesores": sorted({p for f in filas for p in f["profesores"] if p}),
        "asistentesPat": sorted({f["asistentePat"] for f in filas if f["asistentePat"]}),
        "cambios": [],
    }
    contenido = json.dumps(datos["clases"], sort_keys=True, ensure_ascii=False)
    datos["meta"]["hash"] = hashlib.sha256(contenido.encode()).hexdigest()[:16]
    datos["meta"]["generadoEn"] = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    # Advertencias de calidad (no detienen la generación)
    for f in filas:
        if f["dia"] != f["diaCalculado"]:
            print(f"AVISO: {f['clase']} dice '{f['dia']}' pero {f['fecha']} es {f['diaCalculado']}")
        if not f["linkTeams"].startswith("https://"):
            print(f"AVISO: {f['clase']} sin link de Teams válido")
    json.dump(datos, open(salida, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    print(f"OK: {len(filas)} clases, {datos['meta']['horasTotales']} horas, {len(unidades)} unidades, {datos['meta']['inicio']} a {datos['meta']['fin']}")

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
